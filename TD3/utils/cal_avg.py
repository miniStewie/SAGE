from openpyxl import load_workbook


def calculate_column_average(xlsx_path, sheet_name, field_name):
    """
    计算指定 xlsx 文件中，指定 sheet 的指定字段平均值。

    参数：
        xlsx_path: Excel 文件路径
        sheet_name: sheet 名称
        field_name: 字段名，即第一行中的列名

    返回：
        平均值
    """

    # 以只读模式打开，适合较大的 Excel 文件
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)

    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"找不到 sheet：{sheet_name}")

    ws = wb[sheet_name]

    # 读取首行字段名
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))

    if field_name not in header_row:
        wb.close()
        raise ValueError(f"找不到字段：{field_name}，当前字段有：{header_row}")

    # 找到字段所在列，从 1 开始
    col_index = header_row.index(field_name) + 1

    values = []

    # 从第二行开始读取数据
    for row in ws.iter_rows(min_row=2, min_col=col_index, max_col=col_index, values_only=True):
        value = row[0]

        # 跳过空值
        if value is None:
            continue

        # 只统计数字
        if isinstance(value, (int, float)):
            values.append(value)
        else:
            try:
                values.append(float(value))
            except ValueError:
                print(f"跳过非数字值：{value}")

    wb.close()

    if not values:
        raise ValueError(f"字段 {field_name} 中没有可用于计算平均值的数字")

    return sum(values) / len(values)


# ================= 使用示例 =================

xlsx_path = r"C:\Adisk\project\PyCode\BaseEnv\RFL\SAGE\TD3\100info\info29.xlsx"
sheet_name = "step"
field_name = "visible_num"

average_value = calculate_column_average(xlsx_path, sheet_name, field_name)

print(f"{field_name} 的平均值为：{average_value}")